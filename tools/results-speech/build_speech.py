"""Build ABOT-results-speech.docx: the Results and Analysis part of the talk.

Run: python -E -P build_speech.py <pptNprep folder>
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(sys.argv[1])
FIG = ROOT / "assets" / "figures"
OUT = ROOT / "ABOT-results-speech.docx"

# ---- palette (matches the chart PNGs: blue = uses live scores) ----
INK = RGBColor(0x1F, 0x23, 0x28)
NAVY = RGBColor(0x16, 0x32, 0x4F)
ACCENT = "2A78D6"
ACCENT_RGB = RGBColor(0x2A, 0x78, 0xD6)
MUTED = RGBColor(0x5B, 0x65, 0x70)
CUE_FILL = "EEF3F8"
AMBER_FILL = "FFF4E0"
AMBER_LINE = "D98A00"
AMBER_INK = RGBColor(0x7A, 0x4B, 0x00)
GRID_LINE = "C9D3DD"
HEAD_FILL = "DDE7F2"

SERIF = "Georgia"
SANS = "Segoe UI"
SANS_SB = "Segoe UI Semibold"

TEXT_WIDTH_CM = 16.6

# ---------------------------------------------------------------- content
# Spoken paragraphs: (text, core). core=True draws the blue bar: those
# paragraphs alone make the 5-minute version. **...** marks words to stress.

OPENING = {
    "title": "Opening",
    "cue": {"text": "Section slide: “Results”, listing the six result titles."},
    "parts": [
        (None, [
            ("Now to our results. Everything I’ll show comes from **simulation**; no real "
             "radios were used. We have **six results**, and I’ll take them from the "
             "strongest claim to the weakest.", True),
            ("One fact runs through all six. In every test, a route-finder, the method that "
             "picks each next hop, builds a route from a starting radio to a safe zone, our "
             "stand-in for a rescue point. Almost every trip arrives. So we compare "
             "route-finders on **how many hops they take**, where a hop is one radio-to-radio "
             "jump. Fewer is better.", True),
        ]),
    ],
}

RESULTS = [
    {
        "n": 1,
        "title": "Using live link scores changes which routes get picked",
        "paper": "Paper V-C",
        "cue": {"image": "chart_a_link_scores.png", "width": 14.5,
                "text": "Chart A: static vs dynamic Dijkstra, average link-quality score."},
        "parts": [
            ("What we tested", [
                ("First, our strongest result. We ran **Dijkstra**, the classic shortest-route "
                 "method a sat-nav uses, in two versions on the same 100 trips. The **static** "
                 "version counts hops. The **dynamic** version also weighs each link by a live "
                 "link-quality score, and re-plans if the route ahead gets worse.", True),
                ("A link-quality score is a number from 0 to 1 that each radio keeps for how well "
                 "its links have worked recently. We average it over the radios a route passes "
                 "through. Read it as a **relative score** for comparing route-finders, not as "
                 "the share of messages delivered.", True),
            ]),
            ("What happened", [
                ("Look at the left pair of bars. Without live scores, the average link-quality "
                 "score is **0.145**. With them, it’s **0.267**. That’s about **1.8 "
                 "times higher**. Across the damage levels, the gap is about 1.8 to 1.9 times from "
                 "intact to heavy, and about 1.5 times at severe and extreme.", True),
                ("Is that just because the routes differ in length? The dynamic version does take "
                 "longer routes: 3.23 hops on average, against 2.17. But longer routes tend to "
                 "score lower, not higher. If we compare only trips of the same length, the gap "
                 "**widens**, to between 1.9 and 2.6 times.", False),
                ("The right pair is a **supporting check**. In 43 trips, both versions planned "
                 "the same first route, so only re-planning differs. There, the scores are 0.2067 "
                 "against 0.3098. That’s **1.5 times higher**, with a p-value of 0.0057. A "
                 "p-value is the chance of seeing a gap this big if there were really no "
                 "difference. So chance is very unlikely to explain it.", False),
                ("We also tested the A-star pair, a close cousin of Dijkstra, the same way. It "
                 "re-planned in only 3 of 100 trips, so it shows almost no gap, 1.03 times, and "
                 "tells us nothing either way.", False),
            ]),
            ("Why it matters", [
                ("A route-finder that listens to live link quality ends up on better-scoring "
                 "links than one that only counts hops. That’s the result the paper stands "
                 "on. But our claim is **narrow**. The dynamic route-finder chooses its routes "
                 "using this same score, so a higher score is partly expected. So we claim that "
                 "live link-quality scores **change which routes get picked**. We don’t "
                 "claim that more messages arrive.", True),
            ]),
        ],
        "careful": ("Say “link-quality score”, never “delivery rate” or "
                    "“more messages delivered”. Call the 43-trip result a supporting "
                    "check, and don’t put “250 episodes per preset” next to 0.267 "
                    "vs 0.145."),
    },
    {
        "n": 2,
        "title": "The AI route-finder is in the same range as the classic methods",
        "paper": "Paper V-A, Table II",
        "cue": {"image": "chart_b_avg_hops.png", "width": 14.5,
                "text": "Chart B: average hops by route-finder, beside the simplified Table II.",
                "table": {
                    "cols": [2.0, 6.6, 3.6, 2.6],
                    "rows": [
                        ["Grid", "Compared with", "Result", "Adjusted p"],
                        ["10×10", "Dijkstra dynamic", "PPO better", "0.000025"],
                        ["10×10", "A* dynamic, A* static, Dijkstra static",
                         "PPO slightly worse", "0.047"],
                        ["15×15", "Dijkstra dynamic", "PPO better", "0.0019"],
                        ["15×15", "A* dynamic, A* static, Dijkstra static",
                         "No difference", "0.255"],
                    ]}},
        "parts": [
            ("What we tested", [
                ("Second, our AI route-finder, trained with **PPO**, a trial-and-error method "
                 "that learns from points for good moves and penalties for bad ones. We compared "
                 "it with the four classic route-finders, Dijkstra and A-star, each static and "
                 "dynamic, on two grid sizes, 100 trips each.", True),
                ("All five ran at the default damage level, heavy. The chart shows average hops. "
                 "We use averages because hop counts are whole numbers, so the middle values "
                 "often tie.", False),
            ]),
            ("What happened", [
                ("On the 10-by-10 grid, the three fastest classic methods, A-star static, A-star "
                 "dynamic and Dijkstra static, average **2.17** hops. The AI averages **2.48**, "
                 "and Dijkstra dynamic **3.23**. On the 15-by-15 grid, it’s 3.23 to 3.25 for "
                 "the three fastest, 3.56 for the AI, and 4.40 for Dijkstra dynamic.", False),
                ("In the table, the p-values are adjusted because we made several comparisons at "
                 "once. The AI takes fewer hops than Dijkstra dynamic on both grids, with p-values "
                 "well below the usual 0.05 bar. Against the three fastest, it’s slightly "
                 "worse at 10-by-10, with an adjusted p-value of **0.047**, just under the bar. "
                 "At 15-by-15 there’s no significant difference: p is **0.255**.", False),
            ]),
            ("Why it matters", [
                ("So the AI is **in the same range** as the classic methods: slightly behind "
                 "the fastest at 10-by-10, no significant difference at 15-by-15, and better "
                 "than Dijkstra dynamic at both. It learned a sensible way to route, but on this "
                 "task it doesn’t beat simple shortest-route methods. That’s why we call "
                 "the AI part **exploratory**.", True),
            ]),
        ],
        "careful": ("Don’t say “PPO matches the classic planners”: at 10×10 "
                    "it is measurably behind the fastest. Don’t lean on the paper’s "
                    "“lower bound” argument."),
    },
    {
        "n": 3,
        "title": "Making the damage worse barely changes the outcome",
        "paper": "Paper V-B",
        "cue": {"image": "damage_severity.png", "width": 14.5,
                "text": "Damage-sweep figure. Relabel the left axis and title “link-quality "
                        "score”, and replace the “n=” labels with level names."},
        "parts": [
            ("What we tested", [
                ("Third, we made the damage worse, across five damage levels from intact "
                 "to extreme. More damage makes radio signals "
                 "fade faster with distance, as rubble would. Every route-finder ran **250 trips "
                 "at each level**, in five groups of 50.", False),
            ]),
            ("What happened", [
                ("As damage gets worse, only one line on the left moves: **Dijkstra dynamic**. "
                 "Its score sits near **0.28** from intact to heavy, then drops to **0.218** at "
                 "severe and 0.216 at extreme. That’s a single drop of 23 percent at one "
                 "step, not a gradual slide. Every other route-finder moves by less than 0.02. "
                 "On the right, hop counts stay roughly flat for everyone.", True),
                ("Why at severe? Our radio model, which I’ll show last, explains it. At "
                 "heavy, a link works half the time at about 2 kilometres, beyond our "
                 "1.5-kilometre radio range. At severe, that point moves **inside the range**, to "
                 "about 1 kilometre. So severe is the first level where links within range fail "
                 "often.", False),
                ("The AI doesn’t get worse as damage grows. It reaches a safe zone in "
                 "**99.6 percent** of trips at every level, which is one trip in 250. Each trip "
                 "also earns a points score, called reward, and a normal trip scores just over "
                 "100. Four of the AI’s five groups average **101.4 to 102.2**. The fourth "
                 "averages only about 17 from intact to heavy, and about 23 at severe and "
                 "extreme.", False),
            ]),
            ("Why it matters", [
                ("The AI’s one weak spot is a **livelock**, a stuck loop. In one test "
                 "scenario, it keeps picking a neighbour that has died. Each try is rejected "
                 "with a penalty, the situation looks almost the same, so it picks that "
                 "neighbour again. The cause is its **fixed list of eight neighbours**, where "
                 "dead ones stay listed. It happens at every damage level, including intact. So "
                 "the weak spot is how the AI’s choices are set up, **not the damage**.", True),
                ("So in our setup, damage weakens the long links, but it doesn’t stop trips. "
                 "Success stays near 100 percent. This sweep compares route-finders under rising "
                 "stress. It **isn’t proof** that the network is robust to damage.", False),
            ]),
        ],
        "careful": ("Don’t repeat the abstract’s “limitations of RL-based routing "
                    "under extreme network collapse”. Near-100% success is not evidence of "
                    "robustness. The figure on file still says “Packet delivery ratio”: "
                    "fix it, or say aloud that it shows the link-quality score."),
    },
    {
        "n": 4,
        "title": "Our fast simulator makes AI training practical",
        "paper": "Paper IV",
        "cue": {"big": "Up to 7.3× faster than our own reference simulator",
                "text": "One line on the slide. Keep fig_scalability as a backup slide."},
        "parts": [
            ("What we tested", [
                ("Fourth, the simulator itself. Our first simulator handles each radio "
                 "transmission as a separate event. That’s accurate, but too slow for AI "
                 "training, which needed **200,000 practice steps** per run. So we built a faster "
                 "one that moves all the radios forward together, one time step at a time, using "
                 "the same models for radio links, failures and collisions. A collision is when "
                 "two radios send at once and garble each other.", False),
            ]),
            ("What happened", [
                ("Our fast simulator, with collisions switched on as in every result here, runs "
                 "**up to 7.3 times faster** than our own reference simulator. “Up to” "
                 "matters. That figure is from our smallest grid, and the gain is smaller on "
                 "bigger grids. It’s also measured against our own reference, not against "
                 "standard network simulators in general.", True),
                ("Do the two simulators agree? A statistical check called a KS test says they "
                 "agree on **how many radios stay alive**, with a p-value of 0.98 over 20 matched "
                 "runs. They **disagree on traffic figures**, with p below 0.001, because they "
                 "handle collisions differently.", False),
            ]),
            ("Why it matters", [
                ("The fast simulator makes AI training practical. It’s reliable for which "
                 "radios are alive, which is what the AI learns from. But its absolute traffic "
                 "numbers shouldn’t be read as real-world figures.", False),
            ]),
        ],
        "careful": ("Always “up to”, always “than our own reference "
                    "simulator”. Leave out 303× (collisions off, a setting no result "
                    "uses). If asked about 10×10: about 4.7×, from the project files, "
                    "not the paper."),
    },
    {
        "n": 5,
        "title": "Sharing learned values between radios makes no difference",
        "paper": "Paper IV-A",
        "cue": {"text": "The three-row gossip table.",
                "table": {
                    "cols": [5.0, 3.2, 3.2, 2.4],
                    "rows": [
                        ["Measure", "Gossip on", "Gossip off", "p"],
                        ["Reached a safe zone", "100%", "100%", "–"],
                        ["Average hops", "2.487", "2.497", "0.842"],
                        ["Average reward", "101.42", "101.27", "0.826"],
                    ]}},
        "parts": [
            ("What we tested", [
                ("Fifth, gossip. **Gossip** means radios passing information to their "
                 "neighbours, who pass it on, like a rumour. Here, radios share the AI’s "
                 "estimates of how good each position is. We ran the AI with gossip on and off, "
                 "with three training runs and 300 trips for each setting.", False),
            ]),
            ("What happened", [
                ("Both settings reached a safe zone in every trip. Average hops were **2.487** "
                 "with gossip and **2.497** without, with a p-value of 0.842. Average reward was "
                 "101.42 against 101.27, with a p-value of 0.826. Both are far above 0.05.", False),
            ]),
            ("Why it matters", [
                ("So gossip **made no difference**. Its merge step, where a radio would combine "
                 "the values it receives with its own, isn’t connected yet, and both "
                 "settings already solve this task. This shows the system works without gossip, "
                 "not that gossip helps.", True),
            ]),
        ],
        "careful": ("Don’t say gossip helps, or that its airtime cost was measured. The "
                    "line is “Gossip made no difference; its merge step isn’t "
                    "connected yet.”"),
    },
    {
        "n": 6,
        "title": "The radio model behaves sensibly",
        "paper": "Paper V-D, Fig. 2",
        "cue": {"image": "fig4_pdr_vs_distance.png", "width": 9.0,
                "text": "Fig. 2, captioned “Link success vs distance (channel model)”. "
                        "Drop “timestep backend” from the caption."},
        "parts": [
            ("What we tested", [
                ("Last, the radio model underneath it all. This plot shows how often a "
                 "single link gets a message through as distance grows. That’s called the "
                 "packet delivery ratio, or **PDR**. The distance axis is a **log scale**, so "
                 "each labelled gridline is ten times further. The band between the dashed "
                 "lines, from 400 metres to 1.5 kilometres, is where our mesh operates.", False),
            ]),
            ("What happened", [
                ("Each curve is one damage level, a relative stress setting, and more damage "
                 "shifts the curve left. At heavy, a link works half the time at about **2 "
                 "kilometres**. At severe, about **1 kilometre**. At extreme, about **640 "
                 "metres**. At 1.5 kilometres, links succeed 66 percent of the time at heavy, 29 "
                 "percent at severe and 11 percent at extreme. At 400 metres, the nearest "
                 "neighbours, they stay above 73 percent at every level.", False),
            ]),
            ("Why it matters", [
                ("The radio model behaves the way the standard model predicts: more distance and "
                 "more damage mean more failures. That shows **internal consistency, not "
                 "agreement with real radios**. It also explains the damage result: severe is "
                 "where the halfway point moves inside our radio range.", False),
            ]),
        ],
        "careful": ("Say “internal consistency”, never “validated”. "
                    "Don’t quote the intact level’s 45.5 km."),
    },
]

CLOSING = [
    {
        "title": "Taken together",
        "cue": {"text": "Three takeaways, plus one line: “Not yet measured: whether an SOS "
                        "reaches a rescuer.”"},
        "parts": [
            (None, [
                ("So what does this add up to? First, a route-finder that "
                 "uses live link-quality scores **picks different routes**, over links that "
                 "score about 1.8 times higher. Second, the AI routes sensibly, in the same "
                 "range as the classic methods but not ahead of the fastest. Third, we now have "
                 "a simulator fast enough to train AI "
                 "route-finders on a LoRa mesh.", False),
                ("And what we **don’t claim**. We measured link-quality scores along routes. "
                 "We did not measure whether an SOS actually reaches a rescuer, or how long it "
                 "takes. So we don’t claim that more messages get through. That end-to-end "
                 "test is our next experiment.", True),
            ]),
        ],
    },
    {
        "title": "Limits and next steps",
        "cue": {"text": "The five limits as bullets, then “Next steps”."},
        "parts": [
            (None, [
                ("Finally, the limits we state in the paper. Everything is **simulation**, at up "
                 "to 900 radios, with no hardware measurements. Beacons, the short “I’m "
                 "here” messages each radio sends every 30 seconds, take about **4.4 "
                 "percent** of its airtime. That’s high, and a gentler schedule would shift "
                 "every beacon-based result. And the AI picks from only **8 neighbours**, when "
                 "radios have about 31.", False),
                ("The AI’s training didn’t fully settle either. The part that judges "
                 "how good a situation is scored minus 0.67, 0.18 and 0.71 across three training "
                 "runs, where 1 is perfect and below zero is worse than guessing the average. "
                 "And a mesh only exists if people carry the devices.", False),
                ("We also want to be upfront about one thing. In this version, a hop succeeds "
                 "whenever the next radio is still working. **A weak link never makes a hop "
                 "fail.** That’s why every route-finder arrives. So our next steps are hops "
                 "that can fail on weak links, meshes with radios further apart, and measuring "
                 "whether each SOS reaches a rescuer, and how fast. Then comes firmware for the "
                 "LoRa boards the Meshtastic community already uses, to check the simulator "
                 "against real radios.", True),
            ]),
        ],
    },
]

NUMBERS = [
    ("Opening", "Almost every trip reaches a safe zone", "V-A"),
    ("1", "0.145 vs 0.267, about 1.8×", "V-C"),
    ("1", "1.8–1.9× intact to heavy; about 1.5× severe and extreme", "V-C"),
    ("1", "3.23 vs 2.17 hops; 1.9–2.6× when matched on length", "V-C"),
    ("1", "43 trips, 0.2067 vs 0.3098, 1.5× (1.50×), p = 0.0057", "V-C"),
    ("1", "A-star pair: 3 of 100 trips, 1.03×", "V-C"),
    ("2", "100 trips per route-finder per grid", "V-A, Table II caption"),
    ("2", "2.17 / 2.48 / 3.23 and 3.23–3.25 / 3.56 / 4.40 hops", "V-A"),
    ("2", "Adjusted p 0.047 and 0.255; “well below 0.05” = 2.5×10⁻⁵ "
          "and 1.9×10⁻³", "Table II"),
    ("3", "250 trips per level, five groups of 50", "V-B"),
    ("3", "About 0.28, then 0.218 / 0.216; a 23% drop; others move < 0.02", "V-B"),
    ("3", "Half-success distance about 2 km (heavy), about 1 km (severe, 1.03 km)", "V-D"),
    ("3", "99.6% success (0.996). “One trip in 250” is our arithmetic: "
          "(1 − 0.996) × 250", "V-B"),
    ("3", "Reward 101.4–102.2 in four groups; about 17 (≈16.9) and about 23 (≈22.9) "
          "in the fourth", "V-B"),
    ("3", "Eight neighbours to choose from", "III (action space), VI"),
    ("4", "200,000 training steps per run", "III (RL environment)"),
    ("4", "Up to 7.3× faster with collisions on", "IV"),
    ("4", "20 matched runs; p = 0.98 alive radios; p < 0.001 traffic figures", "IV"),
    ("5", "3 training runs, 300 trips each; 100% both; 2.487 / 2.497, p 0.842; "
          "101.42 / 101.27, p 0.826", "IV-A"),
    ("6", "400 m spacing, 1.5 km range", "III (action space), V-D"),
    ("6", "About 2 km / 1 km / 640 m (637 m) half-success distances", "V-D"),
    ("6", "66% / 29% / 11% at 1.5 km (0.664 / 0.286 / 0.109); above 73% at 400 m", "V-D"),
    ("Limits", "Up to 900 radios", "V (intro), VI"),
    ("Limits", "30 s beacons; about 4.4% airtime", "III (packets), VI"),
    ("Limits", "8 neighbours vs about 31 (30.7)", "III, VI"),
    ("Limits", "−0.67, 0.18, 0.71", "VI"),
]

NOT_IN_PAPER = [
    ("Result 1", "“We average it over the radios a route passes through.”",
     "The paper never says how its “mean link PDR” is calculated. "
     "darwinV2/src/train/runner.py:103-104; src/sim/timestep.py:720-733 (context.md §21)."),
    ("Result 3", "“On the right, hop counts stay roughly flat for everyone.”",
     "The damage-sweep figure is not in the paper. results_rev2/figures/damage_severity, from "
     "results_rev1/damage_severity.csv and results_commaware/commaware_severity_episodes.csv."),
    ("Results 3 and 6", "“Why at severe?” and “It also explains the damage "
     "result”",
     "Every number is in V-D; linking them to V-B’s drop is our reading. "
     "ABOT-results-explained.html, result 3; context.md §15."),
    ("Result 4", "“That figure is from our smallest grid, and the gain is smaller on "
     "bigger grids.”",
     "results_rev1/scalability.csv; results_rev1/logs/scalability.log:91 (context.md §12)."),
    ("Taken together; Limits", "“We did not measure whether an SOS actually reaches a "
     "rescuer.”",
     "The paper doesn’t say this, and its abstract and conclusion suggest the opposite. "
     "Per-trip result files hold only success, hops, reward and the link score "
     "(context.md §11, §16)."),
    ("Limits", "“A hop succeeds whenever the next radio is still working.”",
     "darwinV2/src/nav/env.py:311-313 (context.md §11, §15; triage answer 1)."),
]

LEFT_OUT = [
    "45.5 km at the intact level (triage rule 7). Drop it from any damage table on the slides "
    "too.",
    "303×, which is with collisions switched off, a setting no result uses (rule 6).",
    "The ETSI 1% comparison, a European rule (rule 9).",
    "“Lower bound”, and “the comparison is conservative against PPO” (rule 4).",
    "“Steps identical on all 100”, and “250 episodes per preset” beside "
    "0.267 vs 0.145 (rule 11).",
    "The abstract’s “improved packet delivery” and “extreme network "
    "collapse” lines (rules 2 and 3).",
    "The test setup, which the Design and Implementation section already covers.",
]


# ---------------------------------------------------------------- helpers
def set_fonts(rpr_owner, name):
    """Set every rFonts slot (and drop theme fonts, which would win)."""
    rpr = rpr_owner.get_or_add_rPr() if hasattr(rpr_owner, "get_or_add_rPr") else rpr_owner
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn("w:" + attr), None)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn("w:" + attr), name)


def style_run(run, font=SERIF, size=None, bold=None, italic=None, color=None, spacing=None):
    set_fonts(run._r, font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    if spacing is not None:
        rpr = run._r.get_or_add_rPr()
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(spacing))
        rpr.append(sp)
    return run


def para_format(p, before=0, after=0, line=None, indent_left=None, keep_next=False,
                align=None, page_break_before=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    if indent_left is not None:
        pf.left_indent = Cm(indent_left)
    pf.keep_with_next = keep_next
    if align is not None:
        pf.alignment = align
    if page_break_before:
        pf.page_break_before = True
    return p


def add_marked_text(p, text, font=SERIF, size=12.5, color=INK, bold_color=None):
    """Add text with **bold** markup as separate runs."""
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if not chunk:
            continue
        bold = i % 2 == 1
        style_run(p.add_run(chunk), font=font, size=size, bold=bold,
                  color=(bold_color or color) if bold else color)


def left_bar(p, color=ACCENT, size=18, space=9):
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)
    bdr.append(left)
    ppr.append(bdr)


def bottom_rule(p, color=GRID_LINE, size=6, space=4):
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    el = OxmlElement("w:bottom")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    bdr.append(el)
    ppr.append(bdr)


def cell_shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def cell_borders(cell, **edges):
    """edges: top/left/bottom/right = (size_eighths, color) or None for nil."""
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + edge)
        spec = edges.get(edge)
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec[0]))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), spec[1])
        borders.append(el)
    tcpr.append(borders)


def table_setup(table, widths_cm, margins=(100, 160, 100, 160), borders=None):
    """Fixed layout, explicit widths, cell margins (top, left, bottom, right in dxa)."""
    tbl = table._tbl
    tblpr = tbl.tblPr
    # Word hangs the cell border out into the margin by the left cell margin;
    # indent by that much so borders line up with the text, and give the
    # space back from the last column.
    indent = margins[1]
    widths_cm = list(widths_cm)
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), str(indent))
    ind.set(qn("w:type"), "dxa")
    tblpr.append(ind)
    total = sum(widths_cm)
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(int(total * 567)))
    tblw.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    mar = OxmlElement("w:tblCellMar")
    for edge, val in zip(("top", "left", "bottom", "right"), margins):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)
    tb = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        if borders:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), borders)
        else:
            el.set(qn("w:val"), "nil")
        tb.append(el)
    tblpr.append(tb)
    grid = tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths_cm[i] * 567)))
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[i])
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    # put tblPr children in schema order so Word never complains
    order = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
             "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
             "shd", "tblLayout", "tblCellMar", "tblLook"]
    kids = list(tblpr)
    for k in kids:
        tblpr.remove(k)
    kids.sort(key=lambda e: order.index(e.tag.split("}")[1])
              if e.tag.split("}")[1] in order else len(order))
    for k in kids:
        tblpr.append(k)


def no_split_row(row):
    trpr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trpr.append(cs)


def label(p, text, color=ACCENT_RGB, size=8.5):
    style_run(p.add_run(text.upper()), font=SANS_SB, size=size, color=color, spacing=30)


def data_table(container, spec, font_size=9.5):
    rows = spec["rows"]
    t = container.add_table(rows=len(rows), cols=len(rows[0]))
    table_setup(t, spec["cols"], margins=(60, 110, 60, 110), borders=GRID_LINE)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            p = cell.paragraphs[0]
            para_format(p, 0, 0, line=1.1)
            style_run(p.add_run(val), font=SANS_SB if r == 0 else SANS, size=font_size,
                      color=NAVY if r == 0 else INK)
            if r == 0:
                cell_shade(cell, HEAD_FILL)
            else:
                cell_shade(cell, "FFFFFF")
    return t


def cue_box(doc, cue):
    t = doc.add_table(rows=1, cols=1)
    table_setup(t, [TEXT_WIDTH_CM], margins=(140, 220, 160, 220))
    no_split_row(t.rows[0])
    cell = t.cell(0, 0)
    cell_shade(cell, CUE_FILL)
    cell_borders(cell, left=(24, ACCENT))
    p = cell.paragraphs[0]
    para_format(p, 0, 3, keep_next=True)
    label(p, "On screen")
    p = cell.add_paragraph()
    para_format(p, 0, 4, line=1.15, keep_next=bool(cue.get("image") or cue.get("big")
                                                     or cue.get("table")))
    style_run(p.add_run(cue["text"]), font=SANS, size=10, color=MUTED)
    if cue.get("big"):
        p = cell.add_paragraph()
        para_format(p, 4, 4, keep_next=False)
        style_run(p.add_run(cue["big"]), font=SANS_SB, size=15, color=NAVY)
    if cue.get("image"):
        p = cell.add_paragraph()
        para_format(p, 4, 2, keep_next=bool(cue.get("table")))
        p.add_run().add_picture(str(FIG / cue["image"]), width=Cm(cue["width"]))
    if cue.get("table"):
        data_table(cell, cue["table"])
        # a cell must end in a paragraph; make this one as small as possible
        p = cell.add_paragraph()
        para_format(p, 0, 0)
        p.paragraph_format.line_spacing = Pt(2)
        mark = OxmlElement("w:rPr")
        sz_el = OxmlElement("w:sz")
        sz_el.set(qn("w:val"), "2")
        mark.append(sz_el)
        p._p.get_or_add_pPr().append(mark)
    gap = doc.add_paragraph()
    para_format(gap, 0, 6, line=0.8)


def careful_box(doc, text):
    gap = doc.add_paragraph()
    para_format(gap, 0, 2, line=0.8)
    t = doc.add_table(rows=1, cols=1)
    table_setup(t, [TEXT_WIDTH_CM], margins=(110, 220, 120, 220))
    no_split_row(t.rows[0])
    cell = t.cell(0, 0)
    cell_shade(cell, AMBER_FILL)
    cell_borders(cell, left=(24, AMBER_LINE))
    p = cell.paragraphs[0]
    para_format(p, 0, 0, line=1.2)
    label(p, "Careful  ", color=AMBER_INK)
    style_run(p.add_run(text), font=SANS, size=9.5, color=AMBER_INK)


def spoken(doc, text, core):
    p = doc.add_paragraph()
    para_format(p, 0, 9, line=1.42, indent_left=0.55)
    # never split a paragraph across a page turn while reading aloud
    p.paragraph_format.keep_together = True
    if core:
        left_bar(p)
    add_marked_text(p, text, bold_color=NAVY)
    return p


def part_label(doc, text):
    p = doc.add_paragraph()
    para_format(p, 8, 4, keep_next=True, indent_left=0.55)
    label(p, text, color=MUTED, size=8)


def section_heading(doc, eyebrow, title, paper=None, new_page=True):
    p = doc.add_paragraph()
    para_format(p, 0, 2, keep_next=True, page_break_before=new_page)
    label(p, eyebrow, size=9)
    if paper:
        style_run(p.add_run("   ·   " + paper), font=SANS, size=9, color=MUTED)
    h = doc.add_heading(level=1)
    h.text = ""
    para_format(h, 0, 10, keep_next=True, line=1.05)
    style_run(h.add_run(title), font=SANS_SB, size=19, color=NAVY, bold=False)
    return h


def word_count(text):
    return len(re.sub(r"\*\*", "", text).split())


# ---------------------------------------------------------------- document
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
sec.footer_distance = Cm(1.0)

normal = doc.styles["Normal"]
set_fonts(normal.element.get_or_add_rPr(), SERIF)
normal.font.size = Pt(12.5)
normal.font.color.rgb = INK
h1 = doc.styles["Heading 1"]
set_fonts(h1.element.get_or_add_rPr(), SANS_SB)
h1.font.color.rgb = NAVY
h1.font.size = Pt(19)
h1.font.bold = False
h1.paragraph_format.space_before = Pt(0)

# footer: title left, page number right
fp = sec.footer.paragraphs[0]
# The Footer style already has a centre tab and a right tab at 16.5 cm.
style_run(fp.add_run("ABOT · Results and Analysis · spoken script"), font=SANS,
          size=8.5, color=MUTED)
style_run(fp.add_run("\t\t"), font=SANS, size=8.5, color=MUTED)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
r = OxmlElement("w:r")
rpr = OxmlElement("w:rPr")
set_fonts(rpr, SANS)
sz = OxmlElement("w:sz")
sz.set(qn("w:val"), "17")
rpr.append(sz)
col = OxmlElement("w:color")
col.set(qn("w:val"), "5B6570")
rpr.append(col)
r.append(rpr)
t_el = OxmlElement("w:t")
t_el.text = "1"
r.append(t_el)
fld.append(r)
fp._p.append(fld)

# counts for the cover
all_sections = [OPENING] + RESULTS + CLOSING
full_words = sum(word_count(t) for s in all_sections for _, ps in s["parts"] for t, _ in ps)
core_words = sum(word_count(t) for s in all_sections for _, ps in s["parts"] for t, c in ps if c)
full_min = round(full_words / 150)
core_min = round(core_words / 150)

# ---- cover block
p = doc.add_paragraph()
para_format(p, 0, 4)
label(p, "ABOT · IEEE R10-HTC 2026 · talk script", size=9)
p = doc.add_paragraph()
para_format(p, 0, 6, line=1.0)
style_run(p.add_run("Results and Analysis"), font=SANS_SB, size=28, color=NAVY)
p = doc.add_paragraph()
para_format(p, 0, 14, line=1.25)
style_run(p.add_run("The spoken part that follows Design and Implementation. Six results, "
                    "strongest claim first, then what they add up to and the limits."),
          font=SERIF, size=12.5, color=MUTED, italic=True)
bottom_rule(p)

p = doc.add_paragraph()
para_format(p, 6, 8)
for i, (big, small) in enumerate([(f"{full_words:,}", f"words, about {full_min} minutes"),
                                  (f"{core_words:,}", f"words in the 5-minute version")]):
    if i:
        style_run(p.add_run("        "), font=SANS, size=11)
    style_run(p.add_run(big), font=SANS_SB, size=16, color=ACCENT_RGB)
    style_run(p.add_run(" " + small), font=SANS, size=10.5, color=MUTED)

# legend
p = doc.add_paragraph()
para_format(p, 10, 6, keep_next=True)
label(p, "How to read this script", color=NAVY, size=9)
legend = doc.add_table(rows=4, cols=2)
table_setup(legend, [3.4, TEXT_WIDTH_CM - 3.4], margins=(70, 120, 70, 120))
items = [
    ("bar", "Read aloud. Paragraphs with a blue bar form the 5-minute version: read only "
            "those to cut the section short."),
    ("plain", "Read aloud in the full version only."),
    ("cue", "On screen: what the slide shows. Not read aloud."),
    ("careful", "Careful: wording to avoid. Not read aloud."),
]
for r, (kind, text) in enumerate(items):
    a, b = legend.cell(r, 0), legend.cell(r, 1)
    pa = a.paragraphs[0]
    para_format(pa, 0, 0, line=1.1)
    if kind == "bar":
        para_format(pa, 0, 0, line=1.1, indent_left=0.3)
        left_bar(pa)
        style_run(pa.add_run("Spoken text"), font=SERIF, size=10.5, color=INK)
    elif kind == "plain":
        para_format(pa, 0, 0, line=1.1, indent_left=0.3)
        style_run(pa.add_run("Spoken text"), font=SERIF, size=10.5, color=INK)
    elif kind == "cue":
        cell_shade(a, CUE_FILL)
        cell_borders(a, left=(24, ACCENT))
        label(pa, "On screen", size=8)
    else:
        cell_shade(a, AMBER_FILL)
        cell_borders(a, left=(24, AMBER_LINE))
        label(pa, "Careful", color=AMBER_INK, size=8)
    pb = b.paragraphs[0]
    para_format(pb, 0, 0, line=1.15)
    style_run(pb.add_run(text), font=SANS, size=9.5, color=INK)
p = doc.add_paragraph()
para_format(p, 0, 4)
p = doc.add_paragraph()
para_format(p, 0, 0)
style_run(p.add_run("Words in "), font=SANS, size=9.5, color=MUTED)
style_run(p.add_run("bold"), font=SANS_SB, size=9.5, color=NAVY)
style_run(p.add_run(" are the ones to stress."), font=SANS, size=9.5, color=MUTED)

# 5-minute note
p = doc.add_paragraph()
para_format(p, 16, 6, keep_next=True)
label(p, "Short on time?", color=NAVY, size=9)
p = doc.add_paragraph()
para_format(p, 0, 6, line=1.35)
style_run(p.add_run(
    f"Read only the paragraphs with a blue bar: about {core_words:,} words, roughly "
    f"{core_min} minutes. They cover Results 1 to 5 and every framing rule: simulation only "
    "(Opening), no delivery claim (Result 1, Taken together, Limits), the AI is exploratory "
    "(Result 2), and the stuck loop isn’t caused by damage (Result 3)."),
    font=SERIF, size=11.5, color=INK)
p = doc.add_paragraph()
para_format(p, 0, 6, line=1.35)
style_run(p.add_run(
    "In the short version, skip the right pair of bars on Chart A (or crop the chart to the "
    "left pair), let Chart B’s numbers speak for themselves, and keep the gossip table "
    "and Fig. 2 as backup slides."), font=SERIF, size=11.5, color=INK)

# ---- opening: its slide lists the six results
screens = ["Chart A", "Chart B + Table II", "Damage-sweep figure",
           "One line: “up to 7.3×”", "Gossip table", "Fig. 2"]
OPENING["cue"]["table"] = {
    "cols": [0.9, 9.7, 4.8],
    "rows": [["#", "Result", "Its slide shows"]]
            + [[str(r["n"]), r["title"], screens[i]] for i, r in enumerate(RESULTS)],
}
section_heading(doc, "Opening", "Six results, strongest first", new_page=True)
cue_box(doc, OPENING["cue"])
for _, paras in OPENING["parts"]:
    for text, core in paras:
        spoken(doc, text, core)

# ---- results
for res in RESULTS:
    section_heading(doc, f"Result {res['n']}", res["title"], paper=res["paper"])
    cue_box(doc, res["cue"])
    for part, paras in res["parts"]:
        part_label(doc, part)
        for text, core in paras:
            spoken(doc, text, core)
    careful_box(doc, res["careful"])

# ---- closing sections
for s in CLOSING:
    section_heading(doc, "After the results", s["title"])
    cue_box(doc, s["cue"])
    for _, paras in s["parts"]:
        for text, core in paras:
            spoken(doc, text, core)

# ---- notes
section_heading(doc, "Notes", "Every number, and where it is in the paper", new_page=True)
nt = doc.add_table(rows=len(NUMBERS) + 1, cols=3)
table_setup(nt, [2.4, 10.4, 3.8], margins=(55, 100, 55, 100), borders=GRID_LINE)
for c, head in enumerate(["Where", "Number as spoken", "Paper section"]):
    cell = nt.cell(0, c)
    cell_shade(cell, HEAD_FILL)
    pp = cell.paragraphs[0]
    para_format(pp, 0, 0)
    style_run(pp.add_run(head), font=SANS_SB, size=9, color=NAVY)
for r, (where, num, secn) in enumerate(NUMBERS, start=1):
    for c, v in enumerate([("Result " + where) if where.isdigit() else where, num, secn]):
        cell = nt.cell(r, c)
        pp = cell.paragraphs[0]
        para_format(pp, 0, 0, line=1.1)
        style_run(pp.add_run(v), font=SANS, size=9, color=INK)
p = doc.add_paragraph()
para_format(p, 6, 0, line=1.2)
style_run(p.add_run("Charts A and B are not figures in the paper, but every number on them "
                    "comes from V-C and V-A."), font=SANS, size=9, color=MUTED, italic=True)

section_heading(doc, "Notes", "Sentences that rest on project files, not the paper",
                new_page=True)
p = doc.add_paragraph()
para_format(p, 0, 8, line=1.3)
style_run(p.add_run("These help the audience follow, but none of them is a published result. "
                    "If someone asks, say where they come from."), font=SERIF, size=11.5,
          color=INK)
npt = doc.add_table(rows=len(NOT_IN_PAPER) + 1, cols=3)
table_setup(npt, [2.6, 5.6, 8.4], margins=(60, 100, 60, 100), borders=GRID_LINE)
for c, head in enumerate(["Where", "Sentence", "Source"]):
    cell = npt.cell(0, c)
    cell_shade(cell, HEAD_FILL)
    pp = cell.paragraphs[0]
    para_format(pp, 0, 0)
    style_run(pp.add_run(head), font=SANS_SB, size=9, color=NAVY)
for r, row in enumerate(NOT_IN_PAPER, start=1):
    for c, v in enumerate(row):
        cell = npt.cell(r, c)
        if c == 1:
            cell_shade(cell, AMBER_FILL)
        pp = cell.paragraphs[0]
        para_format(pp, 0, 0, line=1.12)
        style_run(pp.add_run(v), font=SANS, size=9, color=INK)

p = doc.add_paragraph()
para_format(p, 16, 6, keep_next=True)
label(p, "Left out on purpose", color=NAVY, size=9)
for item in LEFT_OUT:
    p = doc.add_paragraph(style="List Bullet")
    para_format(p, 0, 3, line=1.2)
    style_run(p.add_run(item), font=SANS, size=10, color=INK)

# ---- put hand-added properties into schema order
SCHEMA = {
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
            "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
            "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
            "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
            "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
            "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
            "sectPr", "pPrChange"],
    "rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike",
            "dstrike", "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid",
            "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs",
            "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em",
            "lang", "eastAsianLayout", "specVanish", "oMath"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
             "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"],
}


def reorder(root):
    for tag, order in SCHEMA.items():
        for el in root.iter(qn("w:" + tag)):
            kids = list(el)
            local = lambda e: e.tag.split("}")[1]
            if all(local(k) in order for k in kids):
                kids.sort(key=lambda e: order.index(local(e)))
                for k in kids:
                    el.remove(k)
                    el.append(k)


# python-docx's template leaves w:zoom without its required percent
for z in doc.settings.element.iter(qn("w:zoom")):
    z.set(qn("w:percent"), "100")

reorder(doc.element.body)
reorder(sec.footer._element)
for st in doc.styles.element.iter(qn("w:style")):
    reorder(st)

doc.core_properties.title = "ABOT Results and Analysis: spoken script"
doc.core_properties.author = "Aniraj Singh and Azzah Mallick"
doc.save(OUT)
print(OUT, full_words, core_words)
