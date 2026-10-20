#!/usr/bin/env python3
"""Build page 2 of the red-team guide (the Q&A prep plan) and its spreadsheet sheet.

Both come from ABOT-qa-prep-plan.md. To change the content, edit the markdown and run this
again; don't edit the HTML or the sheet by hand.

    python -E -P tools/explained-page/build_plan.py [--artifact PATH] [--xlsx]

Writes ABOT-qa-prep-plan.html next to the markdown. It is a standalone page (styles and script
inline) that links back to page 1, ABOT-red-team-explained.html, in the same folder.

--artifact PATH  also writes the copy published inside the claude.ai Artifact, where page 1 is
                 called index.html. Publish it as the file "ABOT-qa-prep-plan.html" next to
                 page 1's --fragment (see build.py).
--xlsx           rewrites the "Q&A prep review" sheet of ABOT-red-team.xlsx. The other sheets
                 are kept, including the cached results of their formulas.

The script stops if the markdown loses the structure it expects, if its per-point table
disagrees with the guide's at-a-glance table, or if a set's points don't match the
per-point table.
"""
import argparse
import importlib.util
import math
import re
import sys
import tempfile
import zipfile
from copy import copy
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
SRC = ROOT / "ABOT-qa-prep-plan.md"
GUIDE_MD = ROOT / "ABOT-red-team-explained.md"
OUT = ROOT / "ABOT-qa-prep-plan.html"
XLSX = ROOT / "ABOT-red-team.xlsx"
SHEET = "Q&A prep review"
GUIDE_HREF = "ABOT-red-team-explained.html"   # page 1, in the same folder
ARTIFACT_GUIDE_HREF = "index.html"            # page 1 inside the claude.ai Artifact

PAGE_TITLE = "ABOT Q&A prep plan"
DESCRIPTION = ("Page 2 of the ABOT red-team guide: the points chosen for Q&A, reviewed and "
               "regrouped into 9 sets, each with an opening line to learn.")

_spec = importlib.util.spec_from_file_location("guide_build", HERE / "build.py")
G = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(G)
esc, attr = G.esc, G.attr

SECTIONS = ["In short", "Your list", "How the sets were judged", "The 9 suggested sets",
            "Every point, one row", "What's wrong with the current sets", "Things not to say",
            "Other advice", "Good choices to keep"]
HEADERS = {
    "The 9 suggested sets": ["Set", "Name", "Opening line to learn", "Points", "How well to know it", "Note"],
    "Every point, one row": ["ID", "The problem in one line", "Priority", "Hard to defend", "Your set",
                             "Suggested set", "Change", "Why", "Prep level"],
    "What's wrong with the current sets": ["#", "Your choice", "The problem", "Why it matters", "The fix"],
    "Things not to say": ["Point", "Tempting line", "Why not", "Say instead"],
    "Other advice": ["#", "Advice", "Why"],
}
CHANGES = {"Keep": "keep", "Move": "move", "Add": "add", "Fix duplicate": "fix"}
PREPS = {"By heart": "heart", "One line": "line", "Glance": "glance"}
NOT_LISTED = "Not on your list"
LETTERS = "ABCDEFGHI"


def fail(msg):
    sys.exit(f"build_plan.py: {msg}")


# ---------------------------------------------------------------- reading the markdown

def load():
    tree = G.build_tree(SRC.read_text(encoding="utf-8"))
    if len(tree.children) != 1 or tree.children[0].level != 1:
        fail("expected one top-level '# ' title")
    h1 = tree.children[0]
    name, _, subtitle = h1.title.partition(": ")
    meta = {}
    for b in h1.blocks:
        if b[0] == "p":
            for line in b[1]:
                m = re.match(r"^\*\*(\w[\w ]*):\*\* (.+)$", line)
                if m:
                    meta[m.group(1)] = m.group(2)
    for k in ("For", "Date", "Built from"):
        if k not in meta:
            fail(f"header line '{k}:' missing")
    if [c.title for c in h1.children] != SECTIONS:
        fail(f"sections changed: {[c.title for c in h1.children]}")
    secs = {c.title: c.blocks for c in h1.children}

    tables = {}
    for title, header in HEADERS.items():
        tbl = [b for b in secs[title] if b[0] == "table"]
        if len(tbl) != 1 or tbl[0][1][0] != header:
            fail(f"table in '{title}' changed")
        tables[title] = tbl[0][1][1]

    # the guide's at-a-glance table is the reference for each point's line, priority and score
    gtree = G.build_tree(GUIDE_MD.read_text(encoding="utf-8"))
    glance = G.parse_glance(G.section_by_title(gtree.children[0], "All 52 at a glance"))

    points, seen = [], set()
    for pid, line, pri, score, yours, sugg, change, why, prep in tables["Every point, one row"]:
        g = glance.get(pid)
        if not g:
            fail(f"{pid} isn't in the guide")
        if (line, pri, int(score)) != (g["line"], g["pri"], g["score"]):
            fail(f"{pid}: line, priority or score differs from the guide's at-a-glance table")
        if pid in seen:
            fail(f"{pid} listed twice in the per-point table")
        seen.add(pid)
        sets = [s.strip() for s in sugg.split("+")]
        if not sets or any(s not in LETTERS for s in sets):
            fail(f"{pid}: suggested set '{sugg}' not understood")
        if change not in CHANGES or prep not in PREPS:
            fail(f"{pid}: change '{change}' or prep level '{prep}' not understood")
        if (yours == NOT_LISTED) != (change == "Add"):
            fail(f"{pid}: only points not on your list can be 'Add'")
        points.append(dict(id=pid, line=line, pri=pri, score=int(score), yours=yours, sets=sets,
                           change=change, why=why, prep=prep))

    sets = []
    for letter, sname, opening, pts, know, note in tables["The 9 suggested sets"]:
        if not (opening.startswith('"') and opening.endswith('"')):
            fail(f"set {letter}: opening line should be in quotes")
        members = []
        for tok in pts.split(", P-"):
            m = re.fullmatch(r"(?:P-)?(\d\d)(?: \(([^)]*)\))?", tok.strip())
            if not m:
                fail(f"set {letter}: points '{pts}' not understood")
            members.append(("P-" + m.group(1), m.group(2) or ""))
        sets.append(dict(letter=letter, name=sname, opening=opening[1:-1], members=members,
                         know=know, note=note))
    if [s["letter"] for s in sets] != list(LETTERS):
        fail("sets should be A to I, in order")
    by_id = {p["id"]: p for p in points}
    for s in sets:
        want = {p["id"] for p in points if s["letter"] in p["sets"]}
        have = {pid for pid, _ in s["members"]}
        if want != have:
            fail(f"set {s['letter']}: points {sorted(have)} but the per-point table says {sorted(want)}")
        for pid, tag in s["members"]:
            if ("new" in tag) != (by_id[pid]["change"] == "Add"):
                fail(f"set {s['letter']}: {pid} should be marked (new) exactly when it's an 'Add'")

    # your list: count the sets you made, and check it names the same points as the table
    your_blocks = secs["Your list"]
    your_intro = " ".join(G.ptext(b) for b in your_blocks if b[0] == "p")
    bullets = [t for b in your_blocks if b[0] == "ul" for t, _ in b[1]]
    singles = re.findall(r"P-\d\d", bullets[0].split("{")[0]) if bullets else []
    braced = [re.findall(r"P-\d\d", g) for t in bullets for g in re.findall(r"\{([^}]*)\}", t)]
    chosen = {p["id"] for p in points if p["yours"] != NOT_LISTED}
    listed = set(singles) | {pid for g in braced for pid in g}
    if listed != chosen:
        fail(f"'Your list' and the per-point table disagree: {sorted(listed ^ chosen)}")
    n_user_sets = len(singles) + len(braced)
    if f"{len(chosen)} points" not in your_intro or f"{n_user_sets} sets" not in your_intro:
        fail(f"'Your list' should say {len(chosen)} points and {n_user_sets} sets")

    def rows(title):
        return tables[title]

    def bullets_of(title):
        return [t for b in secs[title] if b[0] == "ul" for t, _ in b[1]]

    return dict(name=name, subtitle=subtitle, meta=meta, points=points, sets=sets,
                short=bullets_of("In short"), your_intro=your_intro, your=bullets_of("Your list"),
                judge=" ".join(G.ptext(b) for b in secs["How the sets were judged"] if b[0] == "p"),
                wrong=rows("What's wrong with the current sets"), traps=rows("Things not to say"),
                advice=rows("Other advice"), good=bullets_of("Good choices to keep"),
                n_chosen=len(chosen), n_user_sets=n_user_sets)


# ---------------------------------------------------------------- page

def inline(t, here, guide):
    """Markdown inline -> HTML. P-xx links to its row on this page (or to page 1); 'set X' to the set."""
    codes = []

    def keep_code(m):
        codes.append(m.group(1))
        return f"\x01{len(codes) - 1}\x02"

    t = re.sub(r"`([^`]+)`", keep_code, t)
    t = esc(t)
    t = re.sub(r"\[([^\]]+)\]\(([^)\s]+)\)",
               lambda m: m.group(1) if re.search(r"\.(md|xlsx|py|docx|html)$", m.group(2))
               else f'<a href="{attr(m.group(2))}" rel="noopener" target="_blank">{m.group(1)}</a>', t)
    t = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", t)
    t = re.sub(r"(?<![\w*])\*(?=\S)(.+?)(?<=\S)\*(?![\w*])", r"<em>\1</em>", t)
    t = re.sub(r"\{((?:P-\d\d, )*P-\d\d)\}", r'<span class="brace">{\1}</span>', t)
    t = re.sub(r"\x01(\d+)\x02", lambda m: f"<code>{esc(codes[int(m.group(1))])}</code>", t)
    if "*" in re.sub(r"<[^>]+>", "", t):
        fail(f"unconverted asterisk in: {t[:100]}")

    out, in_a, in_code = [], 0, 0
    for part in re.split(r"(<[^>]+>)", t):
        if part.startswith("<"):
            tag = part[1:].split()[0].rstrip(">").lower() if len(part) > 2 else ""
            in_a += {"a": 1, "/a": -1}.get(tag, 0)
            in_code += {"code": 1, "/code": -1}.get(tag, 0)
            out.append(part)
            continue
        if not (in_a or in_code):
            part = re.sub(r"\bP-\d\d\b", lambda m: (
                f'<a class="pref" href="#row-{m.group(0)}">{m.group(0)}</a>' if m.group(0) in here
                else f'<a class="pref" href="{guide}#{m.group(0)}">{m.group(0)}</a>'), part)
            part = re.sub(r"\b([Ss]ets?) ((?:[A-I], )*[A-I](?: and [A-I])?)\b", lambda m: m.group(1) + " " + re.sub(
                r"[A-I]", lambda x: f'<a href="#set-{x.group(0)}">{x.group(0)}</a>', m.group(2)), part)
        out.append(part)
    return "".join(out)


ARROW_R = ('<svg viewBox="0 0 20 20" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" '
           'stroke-linejoin="round" aria-hidden="true"><path d="M4 10h12M11 5l5 5-5 5"/></svg>')
ARROW_L = ('<svg viewBox="0 0 20 20" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" '
           'stroke-linejoin="round" aria-hidden="true"><path d="M16 10H4M9 5l-5 5 5 5"/></svg>')


def prep_class(text):
    t = text.lower()
    return "heart" if "by heart" in t else "line" if "one line" in t else "glance"


def render(d, guide):
    here = {p["id"] for p in d["points"]}

    def il(t):
        return inline(t, here, guide)

    def chip(f, v, label, n, mark):
        return (f'<button type="button" class="chip" data-f="{f}" data-v="{v}" aria-pressed="false">'
                f'{mark}{label} <span class="c">{n}</span></button>')

    pts, sets = d["points"], d["sets"]
    count = lambda key, val: sum(1 for p in pts if p[key] == val)
    n_add = count("change", "Add")
    n_heart, n_glance = count("prep", "By heart"), count("prep", "Glance")

    # ---- bar and hero
    bar = (f'<header class="bar bar-plan"><a class="pagelink back" href="{guide}">{ARROW_L}'
           '<span class="pl-long">The 52 points</span><span class="pl-short">52 points</span></a>'
           '<p class="bar-name"><span class="pg">Page 2</span><span>Q&amp;A prep plan</span></p></header>')
    nav = [("in-short", "In short"), ("your-list", "Your list"), ("sets", "The 9 sets"),
           ("every-point", "Every point"), ("whats-wrong", "What's wrong"),
           ("not-to-say", "Things not to say"), ("advice", "Other advice"), ("keep", "Good choices")]
    stats = [(len(sets), f"sets, down from {d['n_user_sets']}"),
             (d["n_chosen"], f"points on your list, plus {n_add} to add"),
             (n_heart, "points to know by heart"),
             (n_glance, "points only to glance at")]
    hero = ['<header class="hero plan-hero" id="top"><p class="eyebrow">IEEE R10-HTC 2026 · Talk prep · Page 2 of 2</p>',
            f"<h1>{esc(d['name'])}</h1>",
            f'<p class="subtitle">{esc(d["subtitle"][:1].upper() + d["subtitle"][1:])}</p><dl class="meta">']
    for k in ("For", "Date", "Built from"):
        hero.append(f'<div><dt class="lbl">{esc(k)}</dt><dd>{il(d["meta"][k])}</dd></div>')
    hero.append("</dl><ul class=\"stats\">" + "".join(f"<li><b>{n}</b><span>{esc(t)}</span></li>" for n, t in stats)
                + "</ul>")
    hero.append('<nav class="jump" aria-label="On this page"><p class="lbl">On this page</p><ul>'
                + "".join(f'<li><a href="#{i}">{esc(t)}</a></li>' for i, t in nav) + "</ul></nav></header>")

    sec = []
    sec.append('<section class="sec" id="in-short"><h2>In short</h2><div class="prose"><ul>'
               + "".join(f"<li>{il(t)}</li>" for t in d["short"]) + "</ul></div></section>")
    sec.append('<section class="sec" id="your-list"><h2>Your list</h2><div class="prose">'
               f'<p>{il(d["your_intro"])}</p><ul>' + "".join(f"<li>{il(t)}</li>" for t in d["your"])
               + "</ul></div></section>")

    # ---- the 9 sets
    cards = []
    for s in sets:
        chips_ = []
        for pid, tag in s["members"]:
            cls = " new" if "new" in tag else ""
            small = f"<small>{esc(tag)}</small>" if tag else ""
            chips_.append(f'<li><a class="pt-chip{cls}" href="#row-{pid}">{pid}{small}</a></li>')
        pc = prep_class(s["know"])
        cards.append(
            f'<article class="set" id="set-{s["letter"]}">'
            f'<header class="set-h"><span class="set-l" aria-hidden="true">{s["letter"]}</span>'
            f'<h3><span class="vh">Set {s["letter"]}: </span>{il(s["name"])}</h3>'
            f'<span class="prep"><span class="glyph {pc}" aria-hidden="true"></span>{esc(s["know"])}</span></header>'
            f'<figure class="opening"><figcaption class="lbl">Opening line</figcaption>'
            f'<blockquote><p>“{il(s["opening"])}”</p></blockquote>'
            '<button type="button" class="reveal linkbtn" aria-expanded="false">Show the opening line</button></figure>'
            f'<div class="set-pts"><p class="lbl">Points</p><ul class="pt-chips">{"".join(chips_)}</ul></div>'
            f'<p class="note">{il(s["note"])}</p></article>')
    sec.append('<section class="sec" id="sets"><h2>The 9 suggested sets</h2>'
               f'<div class="judge"><p class="lbl">How the sets were judged</p><p>{il(d["judge"])}</p></div>'
               '<div class="practice"><p class="lbl" id="pr-l">Opening lines</p>'
               '<div class="seg" role="group" aria-labelledby="pr-l">'
               '<button type="button" data-practice-btn="show" aria-pressed="true">Show them</button>'
               '<button type="button" data-practice-btn="test" aria-pressed="false">Hide them to test yourself</button>'
               f'</div></div><div class="sets">{"".join(cards)}</div></section>')

    # ---- every point
    rows = []
    for p in pts:
        pri = G.PRI[p["pri"]]
        to = " + ".join(f'<a href="#set-{x}">Set {x}</a>' for x in p["sets"])
        rows.append(
            f'<li class="pp-row" id="row-{p["id"]}" data-change="{CHANGES[p["change"]]}" '
            f'data-prep="{PREPS[p["prep"]]}" data-set="{" ".join(p["sets"])}"><div class="pp-l">'
            f'<p class="pp-top"><a class="pp-id" href="{guide}#{p["id"]}" title="Open the full entry on page 1">{p["id"]}</a>'
            f'<span class="pri">{G.dot(pri)}{p["pri"]}</span>'
            f'<span class="hd"><span class="vh">Hard to defend: </span>{G.meter(p["score"])}{p["score"]}</span></p>'
            f'<p class="pp-line">{il(p["line"])}</p>'
            f'<p class="pp-move"><span class="from">Your set: {il(p["yours"])}</span><span aria-hidden="true">→</span>'
            f'<span class="vh">Suggested: </span><span class="to">{to}</span></p>'
            f'<p class="pp-tags"><span class="chg {CHANGES[p["change"]]}">{esc(p["change"])}</span>'
            f'<span class="prep"><span class="glyph {PREPS[p["prep"]]}" aria-hidden="true"></span>{esc(p["prep"])}</span></p>'
            f'</div><p class="pp-why">{il(p["why"])}</p></li>')
    tools = ['<div class="idx-tools" id="ppTools">',
             '<div class="ctl"><p class="lbl" id="f-chg">Change</p><div class="chips" role="group" aria-labelledby="f-chg">',
             *[chip("change", v, esc(k), count("change", k), f'<span class="sw {v}" aria-hidden="true"></span>')
               for k, v in CHANGES.items()],
             "</div></div>",
             '<div class="ctl"><p class="lbl" id="f-prep">How well to know it</p><div class="chips" role="group" aria-labelledby="f-prep">',
             *[chip("prep", v, esc(k), count("prep", k), f'<span class="glyph {v}" aria-hidden="true"></span>')
               for k, v in PREPS.items()],
             "</div></div>",
             '<div class="ctl"><p class="lbl" id="f-set">Suggested set</p><div class="chips" role="group" aria-labelledby="f-set">',
             *[chip("set", x, f'<span class="lt">{x}</span>', sum(1 for p in pts if x in p["sets"]), "")
               for x in LETTERS],
             "</div></div></div>"]
    sec.append('<section class="sec" id="every-point"><h2>Every point, one row</h2><div class="prose">'
               f'<p>The {d["n_chosen"]} points on your list, plus {n_add} to add. Each ID opens the full entry on page 1. '
               'Filter by the change, by how well to know the point, or by its suggested set.</p></div>'
               + "".join(tools)
               + '<p class="idx-status"><span id="ppCount" aria-live="polite">Showing all points</span>'
               '<button type="button" class="linkbtn" id="ppClear" hidden>Clear filters</button></p>'
               f'<ol class="pp" id="pp" role="list">{"".join(rows)}</ol>'
               '<p class="idx-empty" id="ppEmpty" hidden>No points match all of these filters. Clear one to see more.</p></section>')

    # ---- what's wrong
    probs = []
    for n, choice, prob, why, fix in d["wrong"]:
        probs.append(f'<li class="prob" id="fix-{n}"><div class="prob-h"><span class="prob-n">{esc(n)}</span>'
                     f'<h3>{il(choice)}</h3></div><dl>'
                     f'<div><dt class="lbl">The problem</dt><dd>{il(prob)}</dd></div>'
                     f'<div><dt class="lbl">Why it matters</dt><dd>{il(why)}</dd></div>'
                     f'<div class="fix"><dt class="lbl">The fix</dt><dd>{il(fix)}</dd></div></dl></li>')
    sec.append('<section class="sec" id="whats-wrong"><h2>What\'s wrong with the current sets</h2>'
               f'<ol class="probs" role="list">{"".join(probs)}</ol></section>')

    # ---- things not to say
    traps = []
    for i, (point, tempt, whynot, say) in enumerate(d["traps"], 1):
        label = "Any point" if point == "Any" else point
        traps.append(f'<div class="trap" id="trap-{i}"><p class="trap-k">{il(label)}</p><div class="swap">'
                     f'<div class="avoid"><p class="lbl">Tempting</p><p>{il(tempt)}</p></div>'
                     f'<div class="say"><p class="lbl">Say instead</p><p>{il(say)}</p></div></div>'
                     f'<p class="note"><strong>Why not:</strong> {il(whynot)}</p></div>')
    sec.append(f'<section class="sec" id="not-to-say"><h2>Things not to say</h2><div class="traps">{"".join(traps)}</div></section>')

    # ---- advice, good choices
    sec.append('<section class="sec" id="advice"><h2>Other advice</h2><ol class="advice">'
               + "".join(f'<li><p class="adv">{il(a)}</p><p class="note">{il(w)}</p></li>' for _, a, w in d["advice"])
               + "</ol></section>")
    sec.append('<section class="sec" id="keep"><h2>Good choices to keep</h2><div class="prose"><ul>'
               + "".join(f"<li>{il(t)}</li>" for t in d["good"]) + "</ul></div></section>")

    foot = (f'<footer class="plan-foot"><p>Page 2 of the ABOT red-team guide. '
            f'<a href="{guide}">Back to page 1: the 52 points, explained</a></p>'
            f'<p>Built from <code>ABOT-qa-prep-plan.md</code>, the same source as the "{esc(SHEET)}" sheet in '
            '<code>ABOT-red-team.xlsx</code>.</p></footer>')
    totop = ('<a class="totop" id="totop" href="#top" hidden><svg viewBox="0 0 20 20" fill="none" stroke="currentColor" '
             'stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M10 16V4M5 9l5-5 5 5"/>'
             '</svg>Top</a>')

    css = (HERE / "page.css").read_text(encoding="utf-8") + (HERE / "plan.css").read_text(encoding="utf-8")
    js = (HERE / "plan.js").read_text(encoding="utf-8")
    body = ('<a class="skip" href="#main">Skip to the plan</a>\n' + bar + '\n<main id="main" class="wrap" tabindex="-1">'
            + "".join(hero) + "\n".join(sec) + foot + "</main>\n" + totop + f"\n<script>\n{js}</script>\n")
    return ('<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n'
            '<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">\n'
            f'<meta name="description" content="{attr(DESCRIPTION)}">\n<title>{esc(PAGE_TITLE)}</title>\n'
            '<link rel="preconnect" href="https://fonts.googleapis.com">\n'
            '<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>\n'
            f'<link rel="stylesheet" href="{attr(G.FONTS)}">\n<style>\n{css}</style>\n</head>\n<body>\n'
            + body + "</body>\n</html>\n")


# ---------------------------------------------------------------- spreadsheet

WIDTHS = {"A": 7, "B": 30, "C": 10, "D": 9, "E": 24, "F": 11, "G": 13, "H": 60, "I": 40}
CHANGE_FILL = {"Keep": ("2B7A4B", "E2F1E7"), "Move": ("8A4500", "FFE5CC"), "Add": ("8A2E6A", "F6E7F0"),
               "Fix duplicate": ("9C0006", "FFC7CE")}


def plain(t):
    t = re.sub(r"\[([^\]]+)\]\(([^)\s]+)\)", r"\1", t)
    return t.replace("**", "").replace("`", "").replace("\\*", "*").replace("*", "")


def write_sheet(d):
    import openpyxl
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    # Formula results cached in the other sheets, so they survive the re-save (openpyxl drops them).
    vals_wb = openpyxl.load_workbook(XLSX, data_only=True)
    wb = openpyxl.load_workbook(XLSX)
    cached = {}
    for ws0 in wb.worksheets:
        if ws0.title == SHEET:
            continue
        cached[ws0.title] = {c.coordinate: vals_wb[ws0.title][c.coordinate].value
                             for row in ws0.iter_rows() for c in row
                             if isinstance(c.value, str) and c.value.startswith("=")}

    ref = wb["Without code access"]  # house style: copy fonts, fills and borders from this sheet
    s_title, s_intro, s_head, s_body = ref["A1"], ref["A2"], ref["A3"], ref["F4"]
    s_pri = {}
    for r in range(4, ref.max_row + 1):
        s_pri.setdefault(ref.cell(r, 2).value, ref.cell(r, 2))
    if SHEET in wb.sheetnames:
        pos = wb.sheetnames.index(SHEET)
        del wb[SHEET]
        ws = wb.create_sheet(SHEET, pos)
    else:
        ws = wb.create_sheet(SHEET)
    ws.sheet_properties.tabColor = "8A2E6A"
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w

    font_body = copy(s_body.font)

    def rich(t, bold_all=False):
        t = plain(t) if "**" not in t else re.sub(r"\[([^\]]+)\]\(([^)\s]+)\)", r"\1", t).replace("`", "")
        if "**" not in t:
            return t
        parts = re.split(r"\*\*(.+?)\*\*", t)
        runs = []
        for i, seg in enumerate(parts):
            if seg:
                runs.append(TextBlock(InlineFont(rFont=font_body.name, sz=font_body.sz, b=(i % 2 == 1) or bold_all), seg))
        return CellRichText(*runs)

    def width(c1, c2):
        return sum(WIDTHS[get_column_letter(c)] for c in range(c1, c2 + 1))

    heights = {}

    def put(r, c1, c2, value, like, *, bold=None, fill=None, color=None, border=True, size=None):
        if c2 > c1:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        text = value if isinstance(value, (int, float)) else str(value)
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            cell.font = copy(like.font)
            cell.alignment = copy(like.alignment)
            cell.fill = copy(like.fill)
            if border:
                cell.border = copy(s_body.border)
        a = ws.cell(r, c1)
        a.value = rich(value) if isinstance(value, str) else value
        f = copy(like.font)
        if bold is not None or color or size:
            f = Font(name=f.name, sz=size or f.sz, b=f.b if bold is None else bold,
                     color=color or (f.color.rgb if f.color is not None and f.color.type == "rgb" else None))
            a.font = f
        if fill:
            for c in range(c1, c2 + 1):
                ws.cell(r, c).fill = PatternFill("solid", fgColor=fill)
        a.alignment = Alignment(wrap_text=True, vertical="top")
        chars = max(width(c1, c2) - 2, 4)  # cautious: a little spare height beats clipped text
        lines = sum(max(1, math.ceil(len(plain(p)) / chars)) for p in str(text).split("\n"))
        pt = (f.sz or 10) * 1.32
        heights[r] = max(heights.get(r, 0), lines * pt + 4)

    def header(r, cols):
        for (c1, c2), label in cols:
            put(r, c1, c2, label, s_head)

    def section(r, title):
        ws.row_dimensions[r].height = 10
        put(r + 1, 1, 9, title, s_title, size=12, border=False)
        return r + 2

    # ---- title and how to read the sheet
    put(1, 1, 9, f"{d['name']}: {d['subtitle']}", s_title, border=False)
    put(2, 1, 9, ("Your list of points to prepare for Q&A, reviewed and regrouped into 9 sets (A to I). "
                  "Change: Keep = fine as it is; Move = belongs in a different set; Add = not on your list, worth adding; "
                  "Fix duplicate = listed twice. Prep level: By heart = learn the answer word for word; One line = know one "
                  "sentence; Glance = read once (it only matters if that detail is on a slide). Built from "
                  "ABOT-qa-prep-plan.md, which also builds page 2 of the ABOT-red-team-explained webpage: edit the markdown "
                  "and rebuild (tools/explained-page/build_plan.py --xlsx), because edits made only here won't reach the "
                  f"webpage. Date: {d['meta']['Date']}."), s_intro, border=False)
    r = 3

    r = section(r, "In short")
    for t in d["short"]:
        put(r, 2, 9, t, s_body, border=False)
        r += 1

    r = section(r, "Counts (they update if you edit the Change or Prep level column below)")
    count_rows = r + 1
    header(r, [((2, 2), "What"), ((3, 3), "Count")])
    count_labels = [("Points reviewed", None), ("On your list", None), ("To add (Change = Add)", "Add"),
                    ("Keep", "Keep"), ("Move", "Move"), ("Fix duplicate", "Fix duplicate"),
                    ("By heart", "By heart"), ("One line", "One line"), ("Glance", "Glance")]
    r += 1 + len(count_labels)

    r = section(r, "Your list")
    put(r, 2, 9, d["your_intro"], s_body, border=False)
    r += 1
    for t in d["your"]:
        put(r, 2, 9, "• " + t, s_body, border=False)
        r += 1

    r = section(r, "How the sets were judged")
    put(r, 2, 9, d["judge"], s_body, border=False)
    r += 1

    r = section(r, "The 9 suggested sets")
    header(r, [((1, 1), "Set"), ((2, 2), "Name"), ((3, 5), "Points"), ((6, 7), "How well to know it"),
               ((8, 8), "Opening line to learn"), ((9, 9), "Note")])
    r += 1
    for s in d["sets"]:
        pts = ", ".join(pid + (f" ({tag})" if tag else "") for pid, tag in s["members"])
        put(r, 1, 1, s["letter"], s_body, bold=True)
        put(r, 2, 2, s["name"], s_body, bold=True)
        put(r, 3, 5, pts, s_body)
        put(r, 6, 7, s["know"], s_body)
        put(r, 8, 8, f'"{s["opening"]}"', s_body, bold=True)
        put(r, 9, 9, s["note"], s_body)
        r += 1

    r = section(r, "Every point, one row (filter with the arrows in the header)")
    head_row = r
    header(r, [((i, i), h) for i, h in enumerate(HEADERS["Every point, one row"], 1)])
    r += 1
    p1 = r
    for p in d["points"]:
        put(r, 1, 1, p["id"], s_body)
        put(r, 2, 2, p["line"], s_body)
        put(r, 3, 3, p["pri"], s_pri[p["pri"]])
        put(r, 4, 4, p["score"], s_body)
        put(r, 5, 5, p["yours"], s_body)
        put(r, 6, 6, " + ".join(p["sets"]), s_body, bold=True)
        fg, bg = CHANGE_FILL[p["change"]]
        put(r, 7, 7, p["change"], s_body, bold=True, color=fg, fill=bg)
        put(r, 8, 8, p["why"], s_body)
        put(r, 9, 9, p["prep"], s_body, bold=p["prep"] == "By heart",
            color="808080" if p["prep"] == "Glance" else None)
        r += 1
    p2 = r - 1
    ws.auto_filter.ref = f"A{head_row}:I{p2}"
    ws.conditional_formatting.add(f"D{p1}:D{p2}", ColorScaleRule(
        start_type="num", start_value=1, start_color="63BE7B", mid_type="num", mid_value=50,
        mid_color="FFEB84", end_type="num", end_value=100, end_color="F8696B"))

    # the counts block, now that the per-point rows are known
    new_vals = {}
    pts = d["points"]
    for i, (label, val) in enumerate(count_labels):
        rr = count_rows + i
        if label == "Points reviewed":
            formula, result = f"=COUNTA(A{p1}:A{p2})", len(pts)
        elif label == "On your list":
            formula, result = f'=COUNTIF(E{p1}:E{p2},"<>{NOT_LISTED}")', d["n_chosen"]
        elif val in CHANGES:
            formula, result = f'=COUNTIF(G{p1}:G{p2},"{val}")', sum(1 for p in pts if p["change"] == val)
        else:
            formula, result = f'=COUNTIF(I{p1}:I{p2},"{val}")', sum(1 for p in pts if p["prep"] == val)
        put(rr, 2, 2, label, s_body)
        put(rr, 3, 3, formula, s_body)
        new_vals[f"C{rr}"] = result

    r = section(r, "What's wrong with the current sets")
    header(r, [((1, 1), "#"), ((2, 2), "Your choice"), ((3, 7), "The problem"), ((8, 8), "Why it matters"),
               ((9, 9), "The fix")])
    r += 1
    for n, choice, prob, why, fix in d["wrong"]:
        put(r, 1, 1, int(n), s_body)
        put(r, 2, 2, choice, s_body, bold=True)
        put(r, 3, 7, prob, s_body)
        put(r, 8, 8, why, s_body)
        put(r, 9, 9, fix, s_body)
        r += 1

    r = section(r, "Things not to say")
    header(r, [((1, 2), "Point"), ((3, 7), "Tempting line"), ((8, 8), "Why not"), ((9, 9), "Say instead")])
    r += 1
    for point, tempt, whynot, say in d["traps"]:
        put(r, 1, 2, "Any point" if point == "Any" else point, s_body, bold=True)
        put(r, 3, 7, tempt, s_body, color="9C0006")
        put(r, 8, 8, whynot, s_body)
        put(r, 9, 9, say, s_body, color="2B7A4B")
        r += 1

    r = section(r, "Other advice")
    header(r, [((1, 1), "#"), ((2, 7), "Advice"), ((8, 9), "Why")])
    r += 1
    for n, adv, why in d["advice"]:
        put(r, 1, 1, int(n), s_body)
        put(r, 2, 7, adv, s_body, bold=True)
        put(r, 8, 9, why, s_body)
        r += 1

    r = section(r, "Good choices to keep")
    for t in d["good"]:
        put(r, 2, 9, "• " + t, s_body, border=False)
        r += 1

    for rr, h in heights.items():
        ws.row_dimensions[rr].height = round(max(h, 15), 1)
    ws.sheet_view.zoomScale = 100

    # open the workbook on the new sheet, with only that tab selected
    for other in wb.worksheets:
        other.sheet_view.tabSelected = other is ws
    wb.active = wb.sheetnames.index(SHEET)

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp) / "out.xlsx"
        wb.save(tmp_path)
        cached[SHEET] = new_vals
        restore_cached(tmp_path, XLSX, cached)
    print(f"wrote sheet '{SHEET}' in {XLSX.name}: rows 1-{r - 1}, per-point table rows {p1}-{p2}")


def restore_cached(src, dst, cached):
    """Put formula results back into the saved file, so viewers that don't recalculate still show them."""
    with zipfile.ZipFile(src) as z:
        names = z.namelist()
        files = {n: z.read(n) for n in names}
    wbx = files["xl/workbook.xml"].decode("utf-8")
    rels = files["xl/_rels/workbook.xml.rels"].decode("utf-8")
    targets = dict(re.findall(r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
    targets.update({k: v for v, k in re.findall(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
    for name, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbx):
        name = name.replace("&amp;", "&")
        vals = cached.get(name)
        if not vals:
            continue
        path = "xl/" + targets[rid].lstrip("/").removeprefix("xl/")
        x = files[path].decode("utf-8")
        for ref_, val in vals.items():
            if val is None:
                continue
            pat = re.compile(r'<c r="%s"([^>]*)>(<f>.*?</f>)(?:<v>[^<]*</v>|<v\s*/>)?</c>' % ref_)
            m = pat.search(x)
            if not m:
                fail(f"formula cell {name}!{ref_} not found after saving")
            attrs = re.sub(r'\s+t="[^"]*"', "", m.group(1))
            if isinstance(val, str):
                attrs += ' t="str"'
            v = val if isinstance(val, str) else (int(val) if float(val).is_integer() else val)
            x = x[:m.start()] + f'<c r="{ref_}"{attrs}>{m.group(2)}<v>{G.esc(str(v))}</v></c>' + x[m.end():]
        files[path] = x.encode("utf-8")
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        for n in names:
            z.writestr(n, files[n])


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--artifact", type=Path, help="also write the copy published inside the claude.ai Artifact")
    ap.add_argument("--xlsx", action="store_true", help=f"rewrite the '{SHEET}' sheet of {XLSX.name}")
    args = ap.parse_args()

    d = load()
    page = render(d, GUIDE_HREF)
    OUT.write_text(page, encoding="utf-8", newline="\n")
    print(f"wrote {OUT.relative_to(ROOT)} ({len(page.encode('utf-8')) / 1024:.0f} KB)")
    if args.artifact:
        art = render(d, ARTIFACT_GUIDE_HREF)
        args.artifact.write_text(art, encoding="utf-8", newline="\n")
        print(f"wrote {args.artifact} ({len(art.encode('utf-8')) / 1024:.0f} KB)")
    if args.xlsx:
        write_sheet(d)
    pts = d["points"]
    print(f"points {len(pts)} (on your list {d['n_chosen']}, to add {sum(p['change'] == 'Add' for p in pts)}), "
          f"your sets {d['n_user_sets']}, suggested sets {len(d['sets'])}")
    print("change:", {k: sum(p["change"] == k for p in pts) for k in CHANGES},
          " prep:", {k: sum(p["prep"] == k for p in pts) for k in PREPS})


if __name__ == "__main__":
    main()
